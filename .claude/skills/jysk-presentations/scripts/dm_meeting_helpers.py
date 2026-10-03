#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Slide builders specific to the real DM/SM monthly follow-up meeting
format, pulled directly from Adam's own real decks — see
assets/reference_decks/README.md ("Real 'Домовленості' / AGENDA /
framing patterns") for exactly which files and quotes these come from.
Not generic — these assume the official
assets/official/JYSK_main_template_FY27_indoor.pptx layouts (Agenda =
layout 1, "Заголовок і текст" = layout 21) and their placeholder idx
values, same as every other OLE-table deck built in this project.
"""
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pivot_table_helpers import set_cell, HEADER_FILL, LABEL_FILL

FONT = "Verdana"
BODY_GRAY = RGBColor(0x56, 0x56, 0x55)
BLACK = RGBColor(0x00, 0x00, 0x00)
NAVY = RGBColor(0x14, 0x3C, 0x8A)

AGENDA_LAYOUT_IDX = 1
HEADER_TEXT_LAYOUT_IDX = 21
AGENDA_DURATION_PH_IDX = 10  # the small left "Click to edit agenda text" box
AGENDA_CONTENT_PH_IDX = 2    # the large right "Content Placeholder" box


def agenda_table_slide(prs, items, layout_idx=AGENDA_LAYOUT_IDX):
    """items: list of (topic, minutes) tuples. A real Тема/Час time-
    budget table — confirmed from Adam's own real meeting decks, which
    use exactly this 2-column table plus a total-duration callout (their
    real example: "ТРИ години" for a 180-minute agenda), not a bare
    bullet list. The content placeholder on this layout is nominally an
    OBJECT/content placeholder but python-pptx's installed version
    doesn't give it `.insert_table()` (it resolves to a plain
    SlidePlaceholder, not a content-specific subclass) — so this reads
    the placeholder's own geometry and adds a free table shape at that
    exact position instead, then leaves the empty placeholder alone
    (an empty placeholder's prompt text is EDIT-MODE-ONLY chrome, never
    rendered in the saved file)."""
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    total_min = sum(m for _, m in items)
    h, m = divmod(total_min, 60)
    dur_text = (f"{h} год" + (f" {m} хв" if m else "")) if h else f"{m} хв"

    dur_ph = slide.placeholders[AGENDA_DURATION_PH_IDX]
    tf = dur_ph.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = "РАЗОМ"
    r0.font.name = FONT
    r0.font.size = Pt(14)
    r0.font.bold = True
    r0.font.color.rgb = BODY_GRAY
    p1 = tf.add_paragraph()
    r1 = p1.add_run()
    r1.text = dur_text
    r1.font.name = FONT
    r1.font.size = Pt(36)
    r1.font.bold = True
    r1.font.color.rgb = BLACK

    content_ph = slide.placeholders[AGENDA_CONTENT_PH_IDX]
    left, top, width, height = content_ph.left, content_ph.top, content_ph.width, content_ph.height
    n_rows = len(items) + 1
    gshape = slide.shapes.add_table(n_rows, 2, left, top, width, height)
    tbl = gshape.table
    tbl.columns[0].width = int(width * 0.78)
    tbl.columns[1].width = width - int(width * 0.78)
    set_cell(tbl.cell(0, 0), "Тема", font_size=13, bold=True, fill=HEADER_FILL, align=PP_ALIGN.LEFT)
    set_cell(tbl.cell(0, 1), "Час", font_size=13, bold=True, fill=HEADER_FILL, align=PP_ALIGN.CENTER)
    for i, (topic, minutes) in enumerate(items, start=1):
        set_cell(tbl.cell(i, 0), topic, font_size=12, fill=LABEL_FILL, align=PP_ALIGN.LEFT)
        set_cell(tbl.cell(i, 1), f"{minutes} хв", font_size=12, fill=None, align=PP_ALIGN.CENTER)
    return slide


def goal_statement_slide(prs, title, statement, bullets, add_bullets_fn, layout_idx=HEADER_TEXT_LAYOUT_IDX):
    """The 'Виторг — основна ціль' framing slide: right after the
    results breaker and before any KPI detail, state outright that
    revenue is THE goal and every other KPI that follows is just an
    indicator/lever for it — pulled verbatim in spirit from Adam's own
    real decks ("Всі решта KPI – як індикатор на що потрібно звернути
    увагу та ВПЛИВАТИ для ВИТОРГУ"), which place this slide in exactly
    this spot rather than leaving the framing implicit.
    `add_bullets_fn(slide, bullets, top)` is the caller's own bullet-box
    builder (e.g. build_full_v7.py's `_add_bullets`, which already knows
    this deck's LOGO_SAFE_Y ceiling) — passed in rather than imported,
    since that function is deck-specific, not part of this module."""
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    slide.placeholders[0].text_frame.text = title
    box = slide.shapes.add_textbox(Emu(590549), Emu(1550000), Emu(10500000), Emu(1100000))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = statement
    r.font.name = FONT
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = NAVY
    add_bullets_fn(slide, bullets, Emu(2900000))
    return slide
