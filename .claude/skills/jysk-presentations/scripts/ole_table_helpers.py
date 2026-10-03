#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generalized version of the OLE-embedded-Excel-table pilot: build a
real, double-click-editable Excel worksheet (not a static pptx table)
for a slide, using python-pptx's own public add_ole_object API (tested
code path, not hand-rolled OOXML) plus openpyxl for the embedded
workbook with REAL conditional-formatting rules, and a PIL-rendered
preview image matching JYSK's confirmed look (Verdana, light-blue label
columns, green/red >=100 threshold on index columns).

IMPORTANT: the >=100/<100 rule is only correct for genuine "Index ...
plan/prev." columns, where 100 is a real baseline (plan or prior year).
Mark a raw percentage that has no such baseline (an acceptance rate, a
share picked within a time window) as 'num', not 'pct' — coloring e.g.
a 94% acceptance rate red because it's "below 100" is simply wrong, not
a judgment call. Stock-adjustment % of value has its own col_type,
'stockadj' (below), because it has a real FY27 target that is itself
negative (-0.25%, from fy27_targets.FY27_TARGETS) — neither '100' nor
'0' is the right threshold for it.

Column widths (both the real xlsx and the PNG preview) are computed
from actual content, not guessed by hand — confirmed necessary after
Adam sent a real PowerPoint screenshot showing uneven, shrunk-to-fit
fonts and cramped spacing from the first hand-guessed-width version.
"""
import io
import math
from pptx.util import Emu
from pptx.enum.shapes import PROG_ID
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from PIL import Image, ImageDraw, ImageFont
from fy27_targets import FY27_TARGETS

STOCK_ADJ_TARGET = FY27_TARGETS["stock_adjustment_pct"]  # -0.25%: see fy27_targets.py

LABEL_FILL = "C3D6EB"
HEADER_FILL = "DBE5F1"
GOOD_FILL = "ABEDA5"
BAD_FILL = "FF988C"
OUTER_BORDER = "143C8A"  # JYSK navy — thick outer frame around each table
POTENTIAL_TEXT = "1B5E20"  # dark green — store called out as leading/potential
ATTENTION_TEXT = "B71C1C"  # dark red — store called out as needing attention
POTENTIAL_MARK = "★ "  # ★
ATTENTION_MARK = "⚠ "  # ⚠

FONT_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

DATA_FONT_SIZE = 10
HEADER_FONT_SIZE = 10


def _col_letter(idx):
    """1-based column index -> Excel column letter."""
    s = ""
    while idx > 0:
        idx, r = divmod(idx - 1, 26)
        s = chr(65 + r) + s
    return s


def _marked_name(values, j, highlight):
    """Column j (0-based) of this row, with the ★/⚠ marker prepended if
    it's the name column (index 1) of a highlighted row. Shared between
    the cell-write pass and the auto-width pass so the computed column
    width actually accounts for the marker's extra characters."""
    v = values[j]
    if j == 1 and isinstance(v, str):
        mark = highlight.get(str(values[0]))
        if mark:
            return (POTENTIAL_MARK if mark == "potential" else ATTENTION_MARK) + v
    return v


def _fmt(v, dec):
    if v is None:
        return "—"
    if isinstance(v, str):
        return v
    return f"{v:,.{dec}f}".replace(",", " ")


def build_xlsx_bytes(sheet_name, headers, rows, col_types, col_decimals=None, highlight=None):
    """rows: list of tuples, each len(headers) values + trailing bool
    (bold, e.g. for a district/network total row). col_types: 'label'/
    'num'/'pct' per column, same convention as pivot_table_helpers.py.
    col_decimals: optional list of decimal places per column (default 1)
    — small-magnitude metrics (e.g. a stock-adjustment % of a few tenths
    of a percent) need 2 to stay distinguishable; don't leave at 1 just
    because that's right for index-style values around 100.
    highlight: optional {row_key: 'potential'|'attention'} keyed by the
    row's first-column value (a site/district code) — marks that row's
    label cells with a colored, prefixed name (★ leading / ⚠ needs
    attention) so the table itself carries the same call-out the slide's
    insight bullets make, not just the prose below it."""
    col_decimals = col_decimals or [1] * len(headers)
    highlight = highlight or {}
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    thin = Side(style="thin", color="B9C6D6")
    thick = Side(style="medium", color=OUTER_BORDER)
    n_rows = len(rows) + 1  # + header
    n_cols = len(headers)

    def _border_for(row_idx, col_idx):
        return Border(
            left=thick if col_idx == 1 else thin,
            right=thick if col_idx == n_cols else thin,
            top=thick if row_idx == 1 else thin,
            bottom=thick if row_idx == n_rows else thin,
        )

    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(name="Verdana", size=HEADER_FONT_SIZE, bold=True)
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _border_for(1, j)
    ws.row_dimensions[1].height = 48

    pct_cols = [j for j, t in enumerate(col_types, start=1) if t == "pct"]
    for i, row in enumerate(rows, start=2):
        is_bold = bool(row[-1]) if isinstance(row[-1], bool) else False
        values = row[:-1] if isinstance(row[-1], bool) else row
        row_key = str(values[0])
        mark = highlight.get(row_key)
        ws.row_dimensions[i].height = 22
        for j, v in enumerate(values, start=1):
            ctype = col_types[j - 1]
            display_v = _marked_name(values, j - 1, highlight) if mark else v
            text_color = "000000"
            if mark and ctype == "label":
                is_bold = True
                text_color = POTENTIAL_TEXT if mark == "potential" else ATTENTION_TEXT
            c = ws.cell(row=i, column=j, value=display_v if display_v is not None else None)
            c.font = Font(name="Verdana", size=DATA_FONT_SIZE, bold=is_bold, color=text_color)
            c.border = _border_for(i, j)
            if ctype == "label":
                c.fill = PatternFill("solid", fgColor=LABEL_FILL)
                c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")
                dec = col_decimals[j - 1]
                c.number_format = f"0.{'0' * dec}" if ctype in ("pct", "stockadj") else f"#,##0.{'0' * dec}"

    last_row = len(rows) + 1
    stockadj_cols = [j for j, t in enumerate(col_types, start=1) if t == "stockadj"]
    green_fill = PatternFill("solid", fgColor=GOOD_FILL)
    red_fill = PatternFill("solid", fgColor=BAD_FILL)
    for j in pct_cols:
        col = _col_letter(j)
        rng = f"{col}2:{col}{last_row}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["100"], fill=green_fill))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["100"], fill=red_fill))
    for j in stockadj_cols:
        col = _col_letter(j)
        rng = f"{col}2:{col}{last_row}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=[str(STOCK_ADJ_TARGET)], fill=green_fill))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=[str(STOCK_ADJ_TARGET)], fill=red_fill))

    # Auto-fit column widths (Excel character-width units) from real
    # content — a label column sized for "Inzhur Park, Brovary" instead
    # of a flat guess is what keeps every row at the same font size.
    for j, ctype in enumerate(col_types, start=1):
        header_words = headers[j - 1].replace("\n", " ").split(" ")
        max_word = max((len(w) for w in header_words), default=1)
        data_vals = [
            _marked_name(row[:-1] if isinstance(row[-1], bool) else row, j - 1, highlight)
            for row in rows
        ]
        max_data = max((len(_fmt(v, col_decimals[j - 1])) for v in data_vals), default=1)
        width = max(max_word, max_data) + 2
        if ctype == "label":
            width = max(width, 8)
        ws.column_dimensions[_col_letter(j)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _text_w(draw, text, font):
    bbox = draw.textbbox((0, 0), text, font=font)
    return bbox[2] - bbox[0]


def _wrap_to_width(draw, text, font, max_w):
    lines = []
    for para in text.split("\n"):
        words = para.split(" ")
        cur = ""
        for w in words:
            trial = (cur + " " + w).strip()
            if _text_w(draw, trial, font) <= max_w or not cur:
                cur = trial
            else:
                lines.append(cur)
                cur = w
        lines.append(cur)
    return lines


def build_preview_png(path, headers, rows, col_types, col_decimals=None, highlight=None):
    """Raster 'closed state' snapshot shown on the slide until
    double-clicked. Column widths are computed from real content (the
    widest data value, and the widest single header word so a header
    never has to break a word mid-wrap) — not passed in by hand, which
    is what produced uneven shrunk-to-fit fonts in an earlier version."""
    col_decimals = col_decimals or [1] * len(headers)
    highlight = highlight or {}
    scale = 3
    pad = 16
    row_h = 54
    line_h = 17

    img_probe = Image.new("RGB", (10, 10))
    draw = ImageDraw.Draw(img_probe)
    f_head = ImageFont.truetype(FONT_BOLD, HEADER_FONT_SIZE * scale + 2)
    f_cell = ImageFont.truetype(FONT_PATH, DATA_FONT_SIZE * scale + 2)
    f_cell_b = ImageFont.truetype(FONT_BOLD, DATA_FONT_SIZE * scale + 2)

    def hexrgb(h):
        return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))

    # --- compute each column's width from its actual content ---
    col_w_px = []
    for j, ctype in enumerate(col_types):
        header_words = headers[j].replace("\n", " ").split(" ")
        max_word_w = max((_text_w(draw, w, f_head) for w in header_words), default=0)
        data_texts = []
        for row in rows:
            values = row[:-1] if isinstance(row[-1], bool) else row
            data_texts.append(_fmt(_marked_name(values, j, highlight), col_decimals[j]))
        max_data_w = max((_text_w(draw, t, f_cell_b) for t in data_texts), default=0)
        # `pad` is in pre-scale units but max_word_w/max_data_w are
        # already scaled (measured with *scale-sized fonts) — scale pad
        # too, or this silently reserves far less margin than intended
        # and clips a long left-aligned label (caught via a real render:
        # "1017DISTR06" ran past its column's right edge at pad*2 alone).
        width = max(max_word_w, max_data_w) + pad * scale * 2
        # round UP when converting back to pre-scale units — truncating
        # here previously clipped a long label by a couple of pixels
        # (e.g. "1017DISTR06") once multiplied back out by `scale`.
        col_w_px.append(math.ceil(width / scale))

    header_lines_per_col = [
        _wrap_to_width(draw, h, f_head, w * scale - pad * scale * 2)
        for h, w in zip(headers, col_w_px)
    ]
    header_h = max(len(lines) for lines in header_lines_per_col) * line_h + pad

    W = sum(col_w_px) * scale
    H = (header_h + row_h * len(rows)) * scale
    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)

    def draw_row(y, row, bold):
        x = 0
        values = row[:-1] if isinstance(row[-1], bool) else row
        mark = highlight.get(str(values[0]))
        row_bold = bold or bool(mark)
        for j, (w, val) in enumerate(zip(col_w_px, values)):
            wpx = w * scale
            ctype = col_types[j]
            fill = None
            if ctype == "label":
                fill = hexrgb(LABEL_FILL)
            elif ctype == "pct" and val is not None:
                try:
                    fill = hexrgb(GOOD_FILL if float(val) >= 100 else BAD_FILL)
                except (TypeError, ValueError):
                    fill = None
            elif ctype == "stockadj" and val is not None:
                try:
                    fill = hexrgb(GOOD_FILL if float(val) >= STOCK_ADJ_TARGET else BAD_FILL)
                except (TypeError, ValueError):
                    fill = None
            if fill:
                draw.rectangle([x, y, x + wpx, y + row_h * scale], fill=fill)
            draw.rectangle([x, y, x + wpx, y + row_h * scale], outline=(185, 198, 214))
            text = _fmt(_marked_name(values, j, highlight) if mark else val, col_decimals[j])
            fnt = f_cell_b if row_bold else f_cell
            text_color = (20, 20, 20)
            if mark and ctype == "label":
                text_color = hexrgb(POTENTIAL_TEXT if mark == "potential" else ATTENTION_TEXT)
            tw = _text_w(draw, text, fnt)
            bbox = draw.textbbox((0, 0), text, font=fnt)
            th = bbox[3] - bbox[1]
            tx = x + pad * scale if ctype == "label" else x + (wpx - tw) / 2
            ty = y + (row_h * scale - th) / 2 - bbox[1]
            draw.text((tx, ty), text, fill=text_color, font=fnt)
            x += wpx
        return y + row_h * scale

    x = 0
    for col_w, lines in zip(col_w_px, header_lines_per_col):
        wpx = col_w * scale
        draw.rectangle([x, 0, x + wpx, header_h * scale], fill=hexrgb(HEADER_FILL), outline=(185, 198, 214))
        total_h = len(lines) * line_h * scale
        start_y = (header_h * scale - total_h) / 2
        for li, line in enumerate(lines):
            tw = _text_w(draw, line, f_head)
            draw.text((x + (wpx - tw) / 2, start_y + li * line_h * scale), line, fill=(20, 20, 20), font=f_head)
        x += wpx

    y = header_h * scale
    for row in rows:
        bold = bool(row[-1]) if isinstance(row[-1], bool) else False
        y = draw_row(y, row, bold)

    # Thick outer frame (JYSK navy) around the whole table, on top of
    # every thin per-cell gridline already drawn — Adam asked for the
    # table's overall boundary to read clearly, not just cell edges.
    border_w = 3 * scale
    navy = hexrgb(OUTER_BORDER)
    for i in range(border_w):
        draw.rectangle([i, i, W - 1 - i, H - 1 - i], outline=navy)

    img.save(path)
    return W, H


def add_ole_table(slide, left, top, max_w_emu, max_h_emu, sheet_name, headers, rows, col_types, png_path, col_decimals=None, highlight=None):
    """Build the embedded xlsx + preview PNG and place it as a real,
    double-click-editable Excel object on the slide, sized to fit inside
    the (max_w_emu, max_h_emu) box while preserving the preview's own
    aspect ratio ("object-fit: contain") — never an independently-guessed
    width+height pair, which either stretches the preview or (with a
    bigger, more readable font and more rows) overflows the slide
    entirely, as happened the first time this used a fixed width alone.
    Column widths are computed from content, not passed in by the caller.
    Returns (graphic_frame, width_emu, height_emu) — use the real placed
    size to position whatever comes next (bullets, another table)."""
    xlsx_buf = build_xlsx_bytes(sheet_name, headers, rows, col_types, col_decimals, highlight)
    w_px, h_px = build_preview_png(png_path, headers, rows, col_types, col_decimals, highlight)
    scale = min(max_w_emu / w_px, max_h_emu / h_px)
    width_emu = Emu(int(w_px * scale))
    height_emu = Emu(int(h_px * scale))
    gframe = slide.shapes.add_ole_object(
        object_file=xlsx_buf,
        prog_id=PROG_ID.XLSX,
        left=left, top=top, width=width_emu, height=height_emu,
        icon_file=png_path, icon_width=width_emu, icon_height=height_emu,
    )
    for ole in gframe._element.findall(
        './/{http://schemas.openxmlformats.org/presentationml/2006/main}oleObj'
    ):
        ole.attrib.pop('showAsIcon', None)
    return gframe, width_emu, height_emu
